from pathlib import Path

from docx import Document

from app.schemas.llm import RewriteResponse
from app.services.alignment_service import alignment_service
from app.services.docx_parser import docx_parser
from app.services.validator import validator


def test_docx_parser_extracts_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Gioi thieu", level=1)
    doc.add_paragraph("Doan van thu nhat.")
    doc.add_paragraph("Doan van thu hai.")
    doc.save(path)

    result = docx_parser.parse(str(path))
    assert len(result["paragraphs"]) == 3
    assert result["paragraphs"][1]["section"] == "Gioi thieu"


def test_alignment_maps_excerpt_to_best_paragraph() -> None:
    paragraphs = [
        {"paragraph_ref": "p_0001", "text": "Doan gioi thieu ve he thong."},
        {"paragraph_ref": "p_0002", "text": "Noi dung phan tich du lieu va danh gia."},
    ]
    excerpts = [{"text": "phan tich du lieu", "similarity_score": 0.8, "risk_type": "CLOSE_PARAPHRASE"}]
    units = alignment_service.generate_units("job-1", paragraphs, excerpts)
    assert len(units) == 1
    assert units[0].paragraph_ref == "p_0002"


def test_validator_accepts_reasonable_rewrite() -> None:
    result = RewriteResponse(
        rewritten_text="Doan van da duoc dien dat lai de ngan gon hon nhung van giu nghia chinh.",
        summary_of_changes="Rut gon cau truc.",
        citation_needed=False,
        confidence=0.8,
    )
    score = validator.validate("Doan van da duoc dien dat lai nhung van giu nghia chinh.", result)
    assert score >= 0.35
